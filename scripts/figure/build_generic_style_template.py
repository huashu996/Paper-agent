#!/usr/bin/env python3
"""Build the five-slide public style template entirely from native shapes.

Requires python-pptx. No source deck, image, private text or external asset is read.
The default output is assets/paper-figure-style-template.pptx in this repository.
"""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

NAVY = "142B43"
TEAL = "147D7A"
PALE_TEAL = "EDF7F4"
GOLD = "B18D3F"
PALE_GOLD = "FBF7EB"
GRAY = "52667A"
RULE = "D7E2E9"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"
WIDTH = 13.333333
HEIGHT = 7.5


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def flat(shape):
    """Override any theme shadow without altering native shape geometry."""
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    return shape


def set_run(run, size=18, color=NAVY, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    # Specify the same family for CJK to avoid locale-dependent substitution.
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = OxmlElement("a:ea")
        rpr.append(ea)
    ea.set("typeface", FONT)


def text(slide, name, x, y, w, h, lines, size=18, color=NAVY,
         bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = line
        set_run(run, size, color, bold)
    return shape


def panel(slide, name, x, y, w, h, color=RULE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    flat(s)
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(WHITE)
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(1.25)
    s.adjustments[0] = 0.06
    return s


def box(slide, name, x, y, w, h, zh, en, *, fill=PALE_TEAL, color=TEAL, size=18):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    flat(s)
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(1.25)
    s.adjustments[0] = 0.08
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.14)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, label in enumerate((zh, en)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.06
        run = p.add_run()
        run.text = label
        set_run(run, size, NAVY, i == 0)
    return s


def connector(slide, name, source, target, start=3, end=1, elbow=False):
    # Rectangle connection sites: top=0, left=1, bottom=2, right=3.
    def anchor(s, site):
        return ((s.left + s.width / 2, s.top),
                (s.left, s.top + s.height / 2),
                (s.left + s.width / 2, s.top + s.height),
                (s.left + s.width, s.top + s.height / 2))[site]
    x1, y1 = anchor(source, start)
    x2, y2 = anchor(target, end)
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    c = slide.shapes.add_connector(kind, int(x1), int(y1), int(x2), int(y2))
    flat(c)
    c.name = name
    c.begin_connect(source, start)
    c.end_connect(target, end)
    c.line.color.rgb = rgb(TEAL)
    c.line.width = Pt(1.25)
    # Native DrawingML triangle arrowhead; geometry remains a native connector.
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    c.line._get_or_add_ln().append(tail)
    return c


def base(prs, title_zh, title_en, route, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    text(s, "slide_title", 0.65, 0.42, 12.03, 0.55,
         [f"{title_zh} / {title_en}"], size=28, bold=True)
    text(s, "route_label", 0.65, 1.04, 12.03, 0.35, [route], size=14, color=GRAY)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.55), Inches(12.03), Inches(0.015))
    flat(line)
    line.name = "title_rule"
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(RULE)
    line.line.fill.background()
    text(s, "generic_example_disclosure", 0.65, 6.97, 11.3, 0.25,
         ["通用示意，非研究成果 / Generic illustration, not research results"], size=11, color=GRAY)
    text(s, "page_number", 12.0, 6.97, 0.68, 0.25, [f"{page:02}"], size=11,
         color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def build(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    cp = prs.core_properties
    cp.title = "Paper Agent: Generic Figure Style Template"
    cp.subject = "Five editable academic diagram layouts"
    cp.author = "Paper Agent contributors"
    cp.last_modified_by = "Paper Agent contributors"
    cp.keywords = "generic, editable, academic diagrams, template"
    cp.comments = "Created from scratch. Generic illustrations only; no research results."
    cp.category = "Public documentation template"
    cp.language = "en-US"
    cp.identifier = "paper-agent-generic-style-template"
    cp.version = "1.0"
    cp.revision = 1
    cp.created = cp.modified = datetime(2026, 1, 1)

    # Slides 1–2: background narratives.
    s = base(prs, "研究问题递进", "Research questions", "叙述背景型 / Background narrative", 1)
    xs = (0.65, 4.87, 9.09)
    labels = (("研究背景", "Background"), ("现有局限", "Limitations"), ("研究目标", "Objective"))
    details = (("已知现象", "Known observations"), ("待解问题", "Open questions"), ("验证范围", "Validation scope"))
    nodes = []
    for i, (x, label, detail) in enumerate(zip(xs, labels, details), 1):
        n = box(s, f"narrative_stage_{i}", x, 2.65, 3.59, 1.3, *label)
        nodes.append(n)
        text(s, f"stage_note_{i}", x, 4.48, 3.59, 0.8, list(detail), size=18, color=GRAY, align=PP_ALIGN.CENTER)
    connector(s, "background_to_limitations", nodes[0], nodes[1])
    connector(s, "limitations_to_objective", nodes[1], nodes[2])

    s = base(prs, "来源与目标", "Sources and objectives", "叙述背景型 / Background narrative", 2)
    sources = [box(s, f"source_{i}", 0.85, y, 2.65, 0.95, f"资料{c}", f"Source {c}")
               for i, (y, c) in enumerate(zip((2.0, 3.5, 5.0), "ABC"), 1)]
    scope = box(s, "scope", 5.35, 3.5, 2.65, 0.95, "研究范围", "Study scope")
    goal = box(s, "objective", 9.85, 3.5, 2.65, 0.95, "研究目标", "Objective", fill=PALE_GOLD, color=GOLD)
    for i, n in enumerate(sources, 1):
        connector(s, f"source_{i}_to_scope", n, scope, elbow=(i != 2))
    connector(s, "scope_to_objective", scope, goal)

    # Slides 3–5: method logic.
    s = base(prs, "单向处理管线", "Processing pipeline", "方法逻辑型 / Method logic", 3)
    labels = (("输入", "Input"), ("预处理", "Preprocess"), ("模块A", "Module A"), ("模块B", "Module B"), ("输出", "Output"))
    nodes = [box(s, f"pipeline_node_{i}", 0.65 + i * 2.45, 3.0, 2.2, 1.15, *label)
             for i, label in enumerate(labels)]
    for i in range(4):
        connector(s, f"pipeline_edge_{i}", nodes[i], nodes[i + 1])
    evaluate = box(s, "evaluation", 10.45, 5.1, 2.2, 0.95, "评估", "Evaluation", fill=PALE_GOLD, color=GOLD)
    connector(s, "output_to_evaluation", nodes[-1], evaluate, start=2, end=0)

    s = base(prs, "多分支汇聚", "Branch fusion", "方法逻辑型 / Method logic", 4)
    inp = box(s, "input", 0.65, 3.3, 1.85, 1.05, "输入", "Input")
    branches = [box(s, f"branch_{c}", 4.15, y, 2.65, 1.05, f"模块{c}", f"Module {c}")
                for y, c in zip((1.85, 3.3, 4.75), "ABC")]
    fusion = box(s, "fusion", 8.45, 3.3, 1.85, 1.05, "融合", "Fusion")
    out = box(s, "output", 10.85, 3.3, 1.85, 1.05, "输出", "Output")
    for i, n in enumerate(branches):
        connector(s, f"input_to_branch_{i}", inp, n, elbow=(i != 1))
        connector(s, f"branch_{i}_to_fusion", n, fusion, elbow=(i != 1))
    connector(s, "fusion_to_output", fusion, out)

    s = base(prs, "局部与全局", "Local and global", "方法逻辑型 / Method logic", 5)
    panel(s, "local_panel", 0.65, 1.95, 5.15, 4.25)
    panel(s, "global_panel", 7.55, 1.95, 5.15, 4.25)
    text(s, "local_heading", 1.0, 2.18, 4.45, 0.4, ["局部 / Local"], size=20, color=NAVY, bold=True)
    text(s, "global_heading", 7.9, 2.18, 4.45, 0.4, ["全局 / Global"], size=20, color=NAVY, bold=True)
    inp = box(s, "local_input", 1.25, 2.95, 3.95, 0.95, "输入", "Input")
    local = box(s, "local_update", 1.25, 4.65, 3.95, 0.95, "局部更新", "Local update")
    glob = box(s, "global_fusion", 8.15, 2.95, 3.95, 0.95, "全局融合", "Global fusion")
    out = box(s, "global_output", 8.15, 4.65, 3.95, 0.95, "输出", "Output")
    connector(s, "input_to_local_update", inp, local, start=2, end=0)
    connector(s, "local_update_to_global_fusion", local, glob, elbow=True)
    connector(s, "global_fusion_to_output", glob, out, start=2, end=0)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).resolve().parents[2] / "assets" / "paper-figure-style-template.pptx")
    build(parser.parse_args().output.resolve())
